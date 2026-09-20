"""End-to-end round-trip: SVG → DrawingML → .docx → reopen via python-docx.

These tests verify that:
  * `svg_to_document` injects a `<w:drawing>` anchor into a real document
    without raising.
  * The resulting .docx is valid OOXML (python-docx can read it back).
  * The PresentationML → WordprocessingML rewrite leaves no `p:` shape
    elements behind, and satisfies the structural rules Word enforces
    (every `wps:wsp` ends with a `wps:bodyPr`, every group carries a
    `wpg:cNvPr`, and only the root group is a `wpg:wgp` — nested ones are
    `wpg:grpSp`). Word refuses to open, or silently fails to draw, a file
    that violates these.
"""
from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from lxml import etree

docx = pytest.importorskip("docx")
from docx import Document  # noqa: E402

from mpl_office._inject import (  # noqa: E402
    NS_A,
    NS_P,
    NS_PIC,
    NS_W,
    NS_WPG,
    NS_WPS,
    build_word_drawing,
    presentationml_to_wordml,
)
from mpl_office.docx import Inches, svg_to_document  # noqa: E402

NS = {"w": NS_W, "wps": NS_WPS, "wpg": NS_WPG, "a": NS_A, "pic": NS_PIC, "p": NS_P}

RECT_SVG = (
    '<svg xmlns="http://www.w3.org/2000/svg" width="200" height="100">'
    '<rect x="10" y="20" width="150" height="60" fill="#336699"/>'
    "</svg>"
)

GROUP_SVG = """<svg xmlns="http://www.w3.org/2000/svg" width="200" height="200">
    <g>
        <rect x="0" y="0" width="50" height="50" fill="#ff0000"/>
        <rect x="60" y="0" width="50" height="50" fill="#00ff00"/>
        <circle cx="100" cy="150" r="25" fill="#0000ff"/>
    </g>
</svg>"""

TEXT_SVG = (
    '<svg xmlns="http://www.w3.org/2000/svg" width="200" height="100">'
    '<text x="10" y="50" font-family="DejaVu Sans" font-size="12" '
    'fill="#112233" font-weight="bold">Hello Word</text>'
    "</svg>"
)


def _document_xml(path: Path) -> etree._Element:
    """Read `word/document.xml` out of a saved .docx as an element tree."""
    with zipfile.ZipFile(path) as z:
        return etree.fromstring(z.read("word/document.xml"))


def _assert_word_structural_rules(root: etree._Element) -> None:
    """Assert the invariants Word rejects the file for violating."""
    body_pr = "{%s}bodyPr" % NS_WPS

    wsps = root.findall(".//wps:wsp", NS)
    assert wsps, "expected at least one wps:wsp shape"
    for wsp in wsps:
        assert (
            wsp.find("wps:bodyPr", NS) is not None
        ), "every wps:wsp needs a wps:bodyPr, even with no text"
        assert (
            wsp[-1].tag == body_pr
        ), "wps:bodyPr must be the last child of wps:wsp"

    for wgp in root.findall(".//wpg:wgp", NS):
        assert (
            wgp.find("wpg:cNvPr", NS) is not None
        ), "every wpg:wgp needs a wpg:cNvPr or Word locks the group"

    # `wpg:wgp` is the root element of a wordprocessingGroup and is valid only
    # as the direct child of `a:graphicData`; nested groups must use
    # `wpg:grpSp`. Word silently drops a subtree rooted at a misplaced `wgp`
    # and renders correctly-sized empty space.
    for drawing in root.findall(".//w:drawing", NS):
        roots = drawing.findall(".//a:graphicData/wpg:wgp", NS)
        assert len(roots) == 1, "a drawing needs exactly one root wpg:wgp"
        assert not roots[0].findall(
            ".//wpg:wgp", NS
        ), "nested groups must be wpg:grpSp, not wpg:wgp"

    for grp_sp in root.findall(".//wpg:grpSp", NS):
        assert (
            grp_sp.find("wpg:cNvPr", NS) is not None
        ), "every wpg:grpSp needs a wpg:cNvPr or Word locks the group"
        # CT_WordprocessingGroup requires this leading sequence.
        order = [etree.QName(c).localname for c in grp_sp[:3]]
        assert order == [
            "cNvPr",
            "cNvGrpSpPr",
            "grpSpPr",
        ], "wpg:grpSp children out of schema order: %s" % order

    # No PresentationML shape vocabulary may survive the rewrite.
    for leftover in ("sp", "grpSp", "pic", "nvPr", "txBody"):
        assert not root.findall(".//p:%s" % leftover, NS), (
            "found leftover PresentationML <p:%s> in Word output" % leftover
        )


def test_rect_injection_opens(artifacts: Path):
    doc = Document()
    svg_to_document(RECT_SVG, doc, width=Inches(6), height=Inches(3), source_dpi=96.0)

    out_path = artifacts / "rect_roundtrip.docx"
    doc.save(out_path)
    assert out_path.exists()

    root = _document_xml(out_path)
    drawings = root.findall(".//w:drawing", NS)
    assert len(drawings) == 1, f"expected one drawing, got {len(drawings)}"
    _assert_word_structural_rules(root)

    # python-docx must be able to reopen the file.
    assert len(Document(out_path).paragraphs) >= 1


def test_inline_extent_matches_requested_box(artifacts: Path):
    doc = Document()
    svg_to_document(RECT_SVG, doc, width=Inches(6), height=Inches(3), source_dpi=96.0)
    out_path = artifacts / "extent_roundtrip.docx"
    doc.save(out_path)

    root = _document_xml(out_path)
    extent = root.find(
        ".//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent"
    )
    assert extent is not None
    assert int(extent.get("cx")) == Inches(6)
    assert int(extent.get("cy")) == Inches(3)


def test_group_becomes_word_group(artifacts: Path):
    doc = Document()
    svg_to_document(GROUP_SVG, doc, width=Inches(4), height=Inches(4), source_dpi=96.0)
    out_path = artifacts / "group_roundtrip.docx"
    doc.save(out_path)

    root = _document_xml(out_path)
    # Exactly one root `wpg:wgp` — it is only valid directly under
    # `a:graphicData` — with the converted SVG <g> nested inside it as a
    # `wpg:grpSp`. Emitting a nested group as `wgp` makes Word drop the
    # subtree and render a blank box.
    assert len(root.findall(".//wpg:wgp", NS)) == 1
    assert len(root.findall(".//wpg:grpSp", NS)) >= 1
    assert len(root.findall(".//wps:wsp", NS)) == 3
    _assert_word_structural_rules(root)


def test_text_survives_as_real_word_text(artifacts: Path):
    doc = Document()
    # At 72 DPI a 12px font is 12pt, which keeps the half-point assertion
    # below readable; the conversion itself is DPI-agnostic.
    svg_to_document(TEXT_SVG, doc, source_dpi=72.0)
    out_path = artifacts / "text_roundtrip.docx"
    doc.save(out_path)

    root = _document_xml(out_path)
    assert len(root.findall(".//wps:txbx", NS)) == 1
    runs = root.findall(".//w:txbxContent//w:t", NS)
    assert [r.text for r in runs] == ["Hello Word"]

    # Run properties should carry over: bold, colour, and 12pt as half-points.
    rpr = root.find(".//w:txbxContent//w:r/w:rPr", NS)
    assert rpr is not None
    assert rpr.find("w:b", NS) is not None
    assert rpr.find("w:color", NS).get("{%s}val" % NS_W) == "112233"
    assert rpr.find("w:sz", NS).get("{%s}val" % NS_W) == "24"
    _assert_word_structural_rules(root)


def test_paragraph_and_table_cell_targets(artifacts: Path):
    """A figure can be placed in an existing paragraph or a table cell."""
    doc = Document()

    para = doc.add_paragraph("caption above")
    svg_to_document(RECT_SVG, para, width=Inches(2), height=Inches(1))

    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    svg_to_document(RECT_SVG, cell, width=Inches(2), height=Inches(1))

    out_path = artifacts / "targets_roundtrip.docx"
    doc.save(out_path)

    root = _document_xml(out_path)
    assert len(root.findall(".//w:drawing", NS)) == 2
    # The figure placed in `para` must live in that same paragraph, alongside
    # its text rather than in a new one.
    reopened = Document(out_path)
    assert reopened.paragraphs[0].text == "caption above"
    assert len(reopened.tables[0].cell(0, 0).paragraphs[0].runs) >= 1


def test_rewrite_drops_presentationml_vocabulary():
    """The rewrite pass must translate, not merely re-tag, the fragment."""
    fragment = (
        '<p:sp><p:nvSpPr><p:cNvPr id="2" name="R1"/><p:cNvSpPr/><p:nvPr/>'
        "</p:nvSpPr><p:spPr>"
        '<a:xfrm><a:off x="0" y="0"/><a:ext cx="100" cy="100"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
        "</p:spPr></p:sp>"
    )
    (shape,) = presentationml_to_wordml(fragment)
    assert shape.tag == "{%s}wsp" % NS_WPS
    assert shape.find("wps:cNvPr", NS) is not None, "cNvPr must adopt the wps namespace"
    # `p:nvPr` has no Word equivalent and must be dropped entirely.
    assert shape.find(".//p:nvPr", NS) is None


def test_group_bounds_drive_child_coordinate_space():
    """The wrapper group maps the shapes' own bounds onto the requested box."""
    fragment = (
        '<p:sp><p:nvSpPr><p:cNvPr id="2" name="R1"/><p:cNvSpPr/><p:nvPr/>'
        "</p:nvSpPr><p:spPr>"
        '<a:xfrm><a:off x="100" y="200"/><a:ext cx="400" cy="300"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
        "</p:spPr></p:sp>"
    )
    drawing = build_word_drawing(fragment, width_emu=914400, height_emu=914400)
    xfrm = drawing.find(".//wpg:grpSpPr/a:xfrm", NS)
    assert xfrm.find("a:chOff", NS).get("x") == "100"
    assert xfrm.find("a:chOff", NS).get("y") == "200"
    assert xfrm.find("a:chExt", NS).get("cx") == "400"
    assert xfrm.find("a:chExt", NS).get("cy") == "300"
    assert xfrm.find("a:ext", NS).get("cx") == "914400"


def test_embedded_image_becomes_media_part(artifacts: Path):
    """Raster content inside a figure must become a real Word media part.

    The Rust core emits `r:embed="__mpl_office_img_N__"` sentinels; the
    Word injector has to register each blob on the document part and rewrite
    those sentinels to the relationship ids python-docx hands back.
    """
    matplotlib = pytest.importorskip("matplotlib")
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    fig, ax = plt.subplots(figsize=(3, 2))
    ax.imshow([[0.0, 0.5], [0.5, 1.0]])
    out_path = artifacts / "image_roundtrip.docx"
    try:
        from mpl_office.docx import fig_to_document

        doc = Document()
        fig_to_document(fig, doc, width=Inches(3))
        doc.save(out_path)
    finally:
        plt.close(fig)

    with zipfile.ZipFile(out_path) as z:
        media = [n for n in z.namelist() if n.startswith("word/media/")]
        document = z.read("word/document.xml").decode("utf-8")
        rels = z.read("word/_rels/document.xml.rels").decode("utf-8")

    assert media, "expected the raster to be stored as a media part"
    assert 'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"' in rels
    assert "__mpl_office_img_" not in document, "image sentinel was never rewritten"
    assert "r:embed=" in document


NS_WP_URI = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"


def _wp(tag: str) -> str:
    return "{%s}%s" % (NS_WP_URI, tag)


def test_floating_anchor_structure(artifacts: Path):
    """`anchor="floating"` emits a wp:anchor with position and wrap."""
    doc = Document()
    doc.add_paragraph("body text " * 40)
    svg_to_document(
        RECT_SVG,
        doc,
        width=Inches(2),
        height=Inches(1),
        anchor="floating",
        position=("right", Inches(0.25)),
        relative_to=("margin", "paragraph"),
        wrap="square",
    )
    out_path = artifacts / "floating_roundtrip.docx"
    doc.save(out_path)

    root = _document_xml(out_path)
    anchors = root.findall(".//" + _wp("anchor"))
    assert len(anchors) == 1
    assert not root.findall(".//" + _wp("inline"))

    anchor = anchors[0]
    assert anchor.get("behindDoc") == "0"
    assert anchor.get("simplePos") == "0"

    pos_h = anchor.find(_wp("positionH"))
    assert pos_h.get("relativeFrom") == "margin"
    assert pos_h.find(_wp("align")).text == "right"

    pos_v = anchor.find(_wp("positionV"))
    assert pos_v.get("relativeFrom") == "paragraph"
    assert pos_v.find(_wp("posOffset")).text == str(Inches(0.25))

    assert anchor.find(_wp("wrapSquare")) is not None
    _assert_word_structural_rules(root)


def test_floating_anchor_child_order(artifacts: Path):
    """`wp:anchor` children follow the sequence the schema mandates.

    Word rejects an anchor whose children are out of order, so this pins
    the ordering rather than merely their presence.
    """
    doc = Document()
    svg_to_document(
        RECT_SVG, doc, width=Inches(2), height=Inches(1),
        anchor="floating", wrap="tight",
    )
    out_path = artifacts / "floating_order.docx"
    doc.save(out_path)

    anchor = _document_xml(out_path).find(".//" + _wp("anchor"))
    order = [etree.QName(child).localname for child in anchor]
    expected = [
        "simplePos", "positionH", "positionV", "extent", "effectExtent",
        "wrapTight", "docPr", "cNvGraphicFramePr", "graphic",
    ]
    assert order == expected, order


def test_floating_watermark_behind_text(artifacts: Path):
    doc = Document()
    svg_to_document(
        RECT_SVG, doc, width=Inches(4), height=Inches(2),
        anchor="floating", position=("center", "center"),
        relative_to=("page", "page"), wrap="none", behind_text=True,
    )
    out_path = artifacts / "floating_watermark.docx"
    doc.save(out_path)

    anchor = _document_xml(out_path).find(".//" + _wp("anchor"))
    assert anchor.get("behindDoc") == "1"
    assert anchor.find(_wp("wrapNone")) is not None
    assert anchor.find(_wp("positionH")).find(_wp("align")).text == "center"
    assert anchor.find(_wp("positionH")).get("relativeFrom") == "page"


def test_mixed_anchors_get_unique_ids(artifacts: Path):
    """Every drawing needs a distinct docPr id and stacking order."""
    doc = Document()
    for anchor in ("floating", "inline", "floating"):
        svg_to_document(RECT_SVG, doc, width=Inches(2), height=Inches(1), anchor=anchor)
    out_path = artifacts / "floating_mixed.docx"
    doc.save(out_path)

    root = _document_xml(out_path)
    ids = [e.get("id") for e in root.findall(".//" + _wp("docPr"))]
    assert len(ids) == 3
    assert len(set(ids)) == 3, f"duplicate docPr ids: {ids}"

    heights = [a.get("relativeHeight") for a in root.findall(".//" + _wp("anchor"))]
    assert len(set(heights)) == len(heights), f"duplicate z-order: {heights}"


@pytest.mark.parametrize(
    ("kwargs", "message"),
    [
        ({"anchor": "sideways"}, "anchor must be"),
        ({"anchor": "floating", "wrap": "diagonal"}, "wrap must be one of"),
        ({"anchor": "floating", "position": ("sideways", 0)}, "position must be"),
    ],
)
def test_invalid_anchor_options_raise(kwargs, message):
    doc = Document()
    with pytest.raises(ValueError, match=message):
        svg_to_document(RECT_SVG, doc, width=Inches(2), height=Inches(1), **kwargs)


# ---------------------------------------------------------------------------
# Regression test for the blank-box bug in 0.0.6's .docx output.
# ---------------------------------------------------------------------------


def test_nested_group_is_grpsp_not_wgp(artifacts: Path):
    """Regression: nested groups emitted as `wpg:wgp` render as a blank box.

    `wpg:wgp` is the root element of a wordprocessingGroup and is only valid
    directly under `a:graphicData`. Word does not recognise it as a child of
    another group, drops the whole subtree, and draws the surviving empty root
    group at the requested size — correctly sized, completely blank.
    """
    doc = Document()
    svg_to_document(GROUP_SVG, doc, width=Inches(4), height=Inches(4), source_dpi=96.0)
    out_path = artifacts / "nested_group_grpsp.docx"
    doc.save(out_path)

    root = _document_xml(out_path)
    graphic_roots = root.findall(".//a:graphicData/wpg:wgp", NS)
    assert len(graphic_roots) == 1, "exactly one root group per drawing"

    # Nothing below the root may be a `wgp`.
    assert not graphic_roots[0].findall(".//wpg:wgp", NS)
    # The SVG's <g> must still be there, as a nested group.
    assert graphic_roots[0].findall(".//wpg:grpSp", NS)

    # The shapes must survive inside the nested group, not just at the root —
    # that is the difference between a drawn figure and an empty box.
    nested = graphic_roots[0].find(".//wpg:grpSp", NS)
    assert len(nested.findall(".//wps:wsp", NS)) == 3

    _assert_word_structural_rules(root)
