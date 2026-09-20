"""Word integration — insert matplotlib figures and arbitrary SVG into
`python-docx` documents as native, editable DrawingML vector shapes.

Word expresses DrawingML with a different element vocabulary than
PowerPoint (`wps:wsp` rather than `p:sp`, and a `w:drawing` anchor rather
than a slide's `p:spTree`), so the Rust core's output goes through a
rewrite pass in :mod:`mpl_office._inject` before injection. The geometry,
fills, strokes and text runs themselves are shared between the two
formats.
"""
from __future__ import annotations

from io import StringIO
from typing import Optional

from . import ConvertOptions, convert_svg_to_drawingml_with_images
from ._inject import append_drawing_to_paragraph

__all__ = [
    "fig_to_document",
    "fig_to_paragraph",
    "svg_to_document",
    "Inches",
    "Emu",
]

EMU_PER_INCH = 914_400
EMU_PER_POINT = 12_700


def Inches(v: float) -> int:
    """Inches → EMU, matching python-docx's convention."""
    return int(round(v * EMU_PER_INCH))


def Emu(v: int) -> int:
    return int(v)


def _render_fig_svg(fig, *, dpi: Optional[float] = None) -> tuple[str, float, float]:
    """Render a matplotlib Figure to an in-memory SVG string.

    Forces ``svg.fonttype='none'`` so text comes through as ``<text>``
    elements rather than glyph paths — this is what makes the resulting
    Word output *editable*.

    Returns ``(svg, width_px, height_px)`` where pixel dimensions are at the
    SVG backend's assumed 72 DPI.
    """
    import matplotlib as mpl

    buf = StringIO()
    # matplotlib's SVG backend always writes at 72 DPI.
    with mpl.rc_context({"svg.fonttype": "none"}):
        fig.savefig(buf, format="svg", bbox_inches=None)
    svg = buf.getvalue()
    w_in, h_in = fig.get_size_inches()
    return svg, w_in * 72.0, h_in * 72.0


def svg_to_document(
    svg: str,
    doc_or_paragraph,
    *,
    width: Optional[int] = None,
    height: Optional[int] = None,
    source_dpi: float = 96.0,
    name: str = "Figure",
    anchor: str = "inline",
    position: tuple = (0, 0),
    relative_to: tuple = ("column", "paragraph"),
    wrap: str = "square",
    behind_text: bool = False,
):
    """Inject an SVG string into a `python-docx` document or paragraph.

    `width` and `height` are in EMU (use :func:`Inches`). When omitted the
    figure keeps its natural size at `source_dpi`.

    `doc_or_paragraph` may be a ``Document`` (a new paragraph is appended)
    or an existing ``Paragraph`` (the drawing is added to it, so you can
    place a figure mid-sentence or inside a table cell).

    `anchor` is ``"inline"`` (in the text flow) or ``"floating"`` (text
    flows around it). See :func:`fig_to_document` for the floating options.
    """
    paragraph = _as_paragraph(doc_or_paragraph)

    # The converted shapes are laid out in their own EMU coordinate space;
    # the Word group's `chOff`/`chExt` then maps that space onto the
    # requested on-page box, so no offset is needed here.
    opts = ConvertOptions(
        source_dpi=source_dpi,
        target_width_emu=width,
        target_height_emu=height,
        offset_x_emu=0,
        offset_y_emu=0,
    )
    xml, images = convert_svg_to_drawingml_with_images(svg, opts)

    if width is None or height is None:
        natural_w, natural_h = _natural_size_emu(svg, source_dpi)
        width = width if width is not None else natural_w
        height = height if height is not None else natural_h

    return append_drawing_to_paragraph(
        paragraph,
        xml,
        images,
        width_emu=width,
        height_emu=height,
        name=name,
        anchor=anchor,
        position=position,
        relative_to=relative_to,
        wrap=wrap,
        behind_text=behind_text,
    )


def fig_to_document(
    fig,
    doc_or_paragraph,
    *,
    width: Optional[int] = None,
    height: Optional[int] = None,
    name: str = "Figure",
    anchor: str = "inline",
    position: tuple = (0, 0),
    relative_to: tuple = ("column", "paragraph"),
    wrap: str = "square",
    behind_text: bool = False,
):
    """Render a matplotlib figure into a Word document as native shapes.

    When `width`/`height` are omitted the figure is placed at its natural
    inch dimensions derived from ``fig.get_size_inches()``. Passing only one
    of the two scales the other to preserve the figure's aspect ratio::

        from docx import Document
        from mpl_office.docx import fig_to_document, Inches

        doc = Document("report_template.docx")
        doc.add_heading("Quarterly revenue", level=2)
        fig_to_document(fig, doc, width=Inches(6))
        doc.save("report.docx")

    By default the figure is *inline*: it sits in the text flow like a very
    large character. Pass ``anchor="floating"`` to position it independently
    and have body text wrap around it::

        fig_to_document(
            fig, doc,
            width=Inches(3),
            anchor="floating",
            position=("right", Inches(0.5)),
            relative_to=("margin", "paragraph"),
            wrap="square",
        )

    Floating options:

    * `position` — ``(horizontal, vertical)``. Each is either an EMU offset
      or an alignment keyword (``"left"``, ``"right"``, ``"center"``,
      ``"top"``, ``"bottom"``, ``"inside"``, ``"outside"``).
    * `relative_to` — what each axis is measured against, e.g.
      ``("margin", "paragraph")`` or ``("page", "page")``.
    * `wrap` — ``"square"``, ``"tight"``, ``"through"``, ``"topAndBottom"``
      or ``"none"``.
    * `behind_text` — place the drawing behind the text, for watermarks.
    """
    svg, w_px, h_px = _render_fig_svg(fig)

    aspect = h_px / w_px if w_px else 1.0
    if width is None and height is None:
        width = Inches(w_px / 72.0)
        height = Inches(h_px / 72.0)
    elif height is None:
        height = int(round(width * aspect))
    elif width is None:
        width = int(round(height / aspect)) if aspect else Inches(w_px / 72.0)

    return svg_to_document(
        svg,
        doc_or_paragraph,
        width=width,
        height=height,
        source_dpi=72.0,
        name=name,
        anchor=anchor,
        position=position,
        relative_to=relative_to,
        wrap=wrap,
        behind_text=behind_text,
    )


def fig_to_paragraph(fig, paragraph, **kwargs):
    """Insert a figure into a specific existing paragraph.

    Thin alias for :func:`fig_to_document` — both accept a paragraph — kept
    for symmetry with :func:`mpl_office.pptx.fig_to_placeholder` and for
    readability at call sites that target a paragraph deliberately::

        placeholder = next(p for p in doc.paragraphs if p.text == "{{CHART}}")
        placeholder.text = ""
        fig_to_paragraph(fig, placeholder, width=Inches(6))
    """
    return fig_to_document(fig, paragraph, **kwargs)


def _as_paragraph(doc_or_paragraph):
    """Accept a Document, Paragraph or table _Cell and return a paragraph.

    A ``Document`` gets a fresh trailing paragraph; a ``_Cell`` reuses its
    first (usually empty) paragraph, which is what a caller placing a figure
    in a table cell expects.
    """
    if hasattr(doc_or_paragraph, "add_run"):  # already a Paragraph
        return doc_or_paragraph
    if hasattr(doc_or_paragraph, "add_paragraph"):  # Document or _Cell
        paragraphs = getattr(doc_or_paragraph, "paragraphs", None)
        # A table cell always starts with one empty paragraph; use it rather
        # than leaving a blank line above the figure.
        if paragraphs is not None and len(paragraphs) == 1 and not paragraphs[0].text:
            if type(doc_or_paragraph).__name__ == "_Cell":
                return paragraphs[0]
        return doc_or_paragraph.add_paragraph()
    raise TypeError(
        "Expected a python-docx Document, Paragraph or table cell, got "
        f"{type(doc_or_paragraph).__name__}"
    )


def _natural_size_emu(svg: str, source_dpi: float) -> tuple[int, int]:
    """Best-effort natural size of an SVG in EMU, for when no box is given.

    Reads the root ``width``/``height`` (falling back to the ``viewBox``) and
    scales from `source_dpi` onto EMU. Returns a 1x1 inch box if the SVG
    declares no usable dimensions.
    """
    import re

    def _num(value: str) -> Optional[float]:
        match = re.match(r"\s*([0-9.+-eE]+)\s*(px|pt|in|mm|cm)?\s*$", value or "")
        if not match:
            return None
        try:
            magnitude = float(match.group(1))
        except ValueError:
            return None
        unit = match.group(2)
        # Normalise to source-DPI pixels.
        factors = {
            "in": source_dpi,
            "pt": source_dpi / 72.0,
            "mm": source_dpi / 25.4,
            "cm": source_dpi / 2.54,
        }
        return magnitude * factors.get(unit, 1.0)

    head = svg[:2048]
    width = height = None
    m = re.search(r'\bwidth\s*=\s*"([^"]+)"', head)
    if m:
        width = _num(m.group(1))
    m = re.search(r'\bheight\s*=\s*"([^"]+)"', head)
    if m:
        height = _num(m.group(1))

    if width is None or height is None:
        m = re.search(r'\bviewBox\s*=\s*"([^"]+)"', head)
        if m:
            parts = re.split(r"[\s,]+", m.group(1).strip())
            if len(parts) == 4:
                try:
                    width = width if width is not None else float(parts[2])
                    height = height if height is not None else float(parts[3])
                except ValueError:
                    pass

    if not width or not height:
        return EMU_PER_INCH, EMU_PER_INCH

    emu_per_px = EMU_PER_INCH / source_dpi
    return int(round(width * emu_per_px)), int(round(height * emu_per_px))
