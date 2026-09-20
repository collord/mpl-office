"""Build a small report .docx with several matplotlib figures.

Demonstrates the ways to place a figure in a Word document: appending to
the document, replacing a placeholder paragraph, filling table cells for a
side-by-side layout, and floating a figure so body text wraps around it.

Run:

    .venv/Scripts/python examples/demo_docx.py

and open the resulting ``demo.docx`` in Word. Every chart is native
DrawingML — right-click one and choose "Group > Ungroup" to edit the
individual bars, lines, and labels.
"""
from __future__ import annotations

from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np

from docx import Document

from mpl_office.docx import Inches, fig_to_document


def _bar_figure():
    fig, ax = plt.subplots(figsize=(8, 4))
    ax.bar(["Q1", "Q2", "Q3", "Q4"], [4.2, 5.1, 6.8, 7.5], color="#2E86AB")
    ax.set_title("Revenue by quarter")
    ax.set_ylabel("$M")
    return fig


def _line_figure():
    fig, ax = plt.subplots(figsize=(4, 3))
    x = np.linspace(0, 10, 100)
    ax.plot(x, np.sin(x), label="sin")
    ax.plot(x, np.cos(x), label="cos")
    ax.legend()
    ax.set_title("Trigonometric")
    return fig


def _heatmap_figure():
    fig, ax = plt.subplots(figsize=(4, 3))
    rng = np.random.default_rng(42)
    im = ax.imshow(rng.standard_normal((20, 20)), cmap="viridis")
    ax.set_title("Heatmap")
    fig.colorbar(im, ax=ax)
    return fig


def build():
    out = Path(__file__).parent / "demo.docx"

    doc = Document()
    doc.add_heading("Quarterly Report", level=0)

    # 1. Append a figure directly to the document.
    doc.add_heading("Revenue", level=1)
    doc.add_paragraph(
        "The chart below is native vector art — no rasterized image."
    )
    fig = _bar_figure()
    fig_to_document(fig, doc, width=Inches(6))  # height follows aspect ratio
    plt.close(fig)

    # 2. Replace a placeholder paragraph, the pattern for real templates.
    doc.add_heading("Signals", level=1)
    doc.add_paragraph("{{CHART}}")
    placeholder = next(p for p in doc.paragraphs if p.text == "{{CHART}}")
    placeholder.text = ""
    fig = _line_figure()
    fig_to_document(fig, placeholder, width=Inches(5))
    plt.close(fig)

    # 3. Two figures side by side in a table, including raster content.
    doc.add_heading("Side by side", level=1)
    table = doc.add_table(rows=1, cols=2)
    for cell, make in zip(table.rows[0].cells, (_line_figure, _heatmap_figure)):
        fig = make()
        fig_to_document(fig, cell, width=Inches(3))
        plt.close(fig)

    # 4. A floating figure the body text wraps around.
    doc.add_heading("Floating", level=1)
    fig = _line_figure()
    fig_to_document(
        fig,
        doc,
        width=Inches(2.5),
        anchor="floating",
        position=("right", Inches(0.2)),
        relative_to=("margin", "paragraph"),
        wrap="square",
    )
    plt.close(fig)
    doc.add_paragraph(
        "This paragraph flows around the floating figure anchored to it. "
        * 12
    )

    doc.save(out)
    print(f"wrote {out}")


if __name__ == "__main__":
    build()
